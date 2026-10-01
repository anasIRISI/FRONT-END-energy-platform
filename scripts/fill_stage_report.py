from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = next(Path(r"C:\Users\hp\Downloads").glob("*AVANCEMENT*INFO*.docx"))
OUTPUT = Path(r"C:\Users\hp\Documents\FRONT-END-energy-platform\deliverables\Etat_avancement_stage_4_info_complete.docx")


def write_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


document = Document(SOURCE)
table = document.tables[0]

# Le nom reste volontairement vide : il n'a pas été communiqué par l'étudiant.
write_cell(
    table.cell(1, 1),
    "Conception et développement d’une plateforme intelligente d’aide à la transition énergétique en Belgique, intégrant une application web, une API métier, un module de simulation et un assistant conversationnel IA.",
)
write_cell(
    table.cell(2, 1),
    "Conception de l’architecture full-stack ; développement du front-end React et des API Spring Boot ; catalogue des solutions énergétiques ; formulaire multi-étapes ; création, consultation et historique des simulations ; espace d’administration ; développement de l’interface chatbot et préparation de son intégration avec les recommandations, les rendez-vous et le système de simulation. Validation de la compilation et des principaux parcours API.",
)
write_cell(
    table.cell(3, 1),
    "L’intégration du chatbot a demandé de gérer le contexte métier de manière sécurisée : historique de conversation, catalogue, formulaires, protection de la clé IA et actions contrôlées pour les rendez-vous ou les simulations. La partie simulation a nécessité de fiabiliser les données saisies, notamment la surface et la consommation, d’associer le bon produit au bon formulaire et d’attendre le traitement IA avant d’afficher un résultat enregistré en base de données.",
)
write_cell(
    table.cell(4, 1),
    "Finaliser les tests de bout en bout du chatbot, des rendez-vous et des simulations. Vérifier la persistance des résultats et la cohérence des recommandations. Préparer la démonstration finale, nettoyer les messages d’erreur éventuels et compléter la documentation technique utilisateur et développeur.",
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
