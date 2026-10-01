from pathlib import Path
from docx import Document

source = next(Path(r"C:\Users\hp\Downloads").glob("*AVANCEMENT*INFO*.docx"))
output = Path(r"C:\Users\hp\Documents\FRONT-END-energy-platform\deliverables\Etat_avancement_stage_4_info_complete.docx")

for name, path in (("SOURCE", source), ("OUTPUT", output)):
    doc = Document(path)
    section = doc.sections[0]
    print(name)
    print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "sections", len(doc.sections))
    print("page", section.page_width, section.page_height, "margins", section.left_margin, section.right_margin, section.top_margin, section.bottom_margin)
    print("rows", len(doc.tables[0].rows), "columns", len(doc.tables[0].columns))
