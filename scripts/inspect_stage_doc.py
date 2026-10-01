import sys
from pathlib import Path
from docx import Document

source = Path(sys.argv[1]) if len(sys.argv) > 1 else next(Path(r"C:\Users\hp\Downloads").glob("*AVANCEMENT*INFO*.docx"))
document = Document(source)

print("PARAGRAPHES")
for index, paragraph in enumerate(document.paragraphs):
    if paragraph.text.strip():
        print(f"{index}: {paragraph.text}")

print(f"TABLES {len(document.tables)}")
for table_index, table in enumerate(document.tables):
    rows = [[cell.text.replace("\n", " | ") for cell in row.cells] for row in table.rows]
    print(f"TABLE {table_index}: {rows}")
